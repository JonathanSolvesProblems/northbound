"""Extract the hackathon workbook's sheets to CSV.

The participant portal hands out one .xlsx with five sheets: Tlorder, Dispatch,
Driver, Trucks, Trailers. Only the Driver sheet is offered separately as a CSV, so
it is easy to mistake it for the whole dataset. It is not: Dispatch carries 10,479
completed legs with an empty/loaded flag and a distance per leg, which is where
every number in the README comes from.

Put the workbook at data/raw/Hackathon_Data.xlsx, then run:

    python scripts/extract_sheets.py
"""
import csv, os, re, sys

try:
    import openpyxl
except ImportError:
    sys.exit('openpyxl is required:  pip install openpyxl')

ROOT = os.path.join(os.path.dirname(__file__), '..')
SRC = os.path.join(ROOT, 'data', 'raw', 'Hackathon_Data.xlsx')
OUT = os.path.join(ROOT, 'data', 'raw', 'sheets')


def main() -> int:
    if not os.path.exists(SRC):
        sys.exit(
            f'Workbook not found at {SRC}\n'
            'Download it from the participant portal (Data Sets) and put it there.'
        )

    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    os.makedirs(OUT, exist_ok=True)
    print(f'sheets: {", ".join(wb.sheetnames)}\n')

    for name in wb.sheetnames:
        rows = wb[name].iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            print(f'  {name}: empty, skipped')
            continue

        header = [str(h) if h is not None else '' for h in header]
        dest = os.path.join(OUT, re.sub(r'[^A-Za-z0-9_]', '_', name).lower() + '.csv')
        written = 0
        with open(dest, 'w', newline='', encoding='utf-8') as fh:
            w = csv.writer(fh)
            w.writerow(header)
            for row in rows:
                if all(c is None for c in row):
                    continue
                w.writerow(['' if c is None else c for c in row])
                written += 1
        print(f'  {name:<10} {written:>7,} rows x {len(header):>3} cols  -> {os.path.relpath(dest, ROOT)}')

    wb.close()
    print('\nNote: 1980-01-01 is their null-date sentinel, and every empty repositioning')
    print('leg carries it in LS_EXPECTED_DATE. Use PICKUP_BY or PLAN_DEPART instead.')
    return 0


if __name__ == '__main__':
    sys.exit(main())
